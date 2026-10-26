#!/usr/bin/env python3
"""Generate READIT's test documents.

Every fixture is written from scratch by this script, so the files are free to
redistribute with the project. Run:  python3 tools/make_fixtures.py
Output: samples/  (PDF, DOCX, PPTX, EPUB, RTF, HTML, Markdown, TXT, plus
deliberately broken files used by the error-handling tests).
"""
import io
import os
import zipfile

from docx import Document
from pptx import Presentation
from pptx.util import Inches, Pt
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

SERIF = "Times-Roman"
for path in ["/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf"]:
    if os.path.exists(path):
        pdfmetrics.registerFont(TTFont("DejaVuSerif", path))
        SERIF = "DejaVuSerif"  # has the cedi sign; Times-Roman does not
from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfReader, PdfWriter
from ebooklib import epub

OUT = os.path.join(os.path.dirname(__file__), "..", "samples")
os.makedirs(OUT, exist_ok=True)

TITLE = "Know Your Customer Basics"
INTRO = (
    "Financial institutions in Ghana must verify the identity of every customer. "
    "Dr. Mensah notes that 2.5 million accounts were reviewed in 2024. "
    "The penalty for a breach can reach GH₵ 1,250.50 per day."
)
RISK = (
    "Customers are rated low, medium or high risk. High-risk customers need enhanced due diligence, "
    "e.g. source-of-funds checks, before the account opens on 01/03/2026."
)
STEPS = ["Collect a valid Ghana Card.", "Screen the name against sanctions lists.", "Record the result in the KYC file."]


def p(name):
    return os.path.join(OUT, name)


def make_docx():
    d = Document()
    d.core_properties.title = TITLE
    d.add_heading(TITLE, level=1)
    d.add_paragraph(INTRO)
    d.add_heading("Risk rating", level=2)
    d.add_paragraph(RISK)
    for s in STEPS:
        d.add_paragraph(s, style="List Number")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Risk", "Review cycle"
    t.cell(1, 0).text, t.cell(1, 1).text = "High", "Every 12 months"
    section = d.sections[0]
    section.header.paragraphs[0].text = "CONFIDENTIAL HEADER"
    section.footer.paragraphs[0].text = "Footer text"
    d.save(p("kyc-basics.docx"))


def make_pptx():
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = TITLE
    s1.placeholders[1].text = "Training for front-office staff"
    s1.notes_slide.notes_text_frame.text = "Welcome everyone and introduce the session."
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Three steps"
    body = s2.placeholders[1].text_frame
    body.text = STEPS[0]
    for s in STEPS[1:]:
        body.add_paragraph().text = s
    s2.notes_slide.notes_text_frame.text = "Stress that screening happens before onboarding."
    s3 = prs.slides.add_slide(prs.slide_layouts[5])
    s3.shapes.title.text = "Review cycles"
    rows, cols = 2, 2
    tbl = s3.shapes.add_table(rows, cols, Inches(1), Inches(2), Inches(6), Inches(1.5)).table
    tbl.cell(0, 0).text, tbl.cell(0, 1).text = "Risk", "Review cycle"
    tbl.cell(1, 0).text, tbl.cell(1, 1).text = "High", "Every 12 months"
    prs.core_properties.title = TITLE
    prs.save(p("kyc-training.pptx"))


def wrap(c, text, x, y, width, font=None, size=11, leading=14):
    font = font or SERIF
    words, line = text.split(), ""
    for w in words:
        trial = (line + " " + w).strip()
        if c.stringWidth(trial, font, size) > width:
            c.drawString(x, y, line)
            y -= leading
            line = w
        else:
            line = trial
    if line:
        c.drawString(x, y, line)
        y -= leading
    return y


def make_pdf():
    c = canvas.Canvas(p("kyc-guide.pdf"), pagesize=A4)
    c.setTitle(TITLE)
    w, h = A4
    pages = [
        [("h1", TITLE), ("p", INTRO), ("h2", "Risk rating"), ("p", RISK)],
        [("h2", "Onboarding steps")] + [("p", f"{i + 1}. {s}") for i, s in enumerate(STEPS)] + [("p", "Each step is logged. " * 12)],
        [("h2", "Record keeping"), ("p", "Records are kept for at least five years after the relationship ends. " * 5)],
    ]
    for n, blocks in enumerate(pages, start=1):
        c.setFont("Helvetica", 8)
        c.drawString(72, h - 40, "e-Crime Bureau Training Notes")
        y = h - 90
        for kind, text in blocks:
            if kind == "h1":
                c.setFont("Helvetica-Bold", 20)
                c.drawString(72, y, text)
                y -= 34
            elif kind == "h2":
                y -= 8
                c.setFont("Helvetica-Bold", 14)
                c.drawString(72, y, text)
                y -= 24
            else:
                c.setFont(SERIF, 11)
                y = wrap(c, text, 72, y, w - 144)
                y -= 10
        c.setFont("Helvetica", 8)
        c.drawCentredString(w / 2, 30, f"Page {n} of {len(pages)}")
        c.showPage()
    c.save()


def make_scanned_pdf():
    img = Image.new("RGB", (1240, 1754), "white")
    d = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 40)
    except OSError:
        font = ImageFont.load_default()
    d.text((120, 200), "SCANNED NOTICE", fill="black", font=font)
    d.text((120, 300), "This page is an image with no text layer.", fill="black", font=font)
    img.save(p("scanned-notice.pdf"), "PDF", resolution=150)
    img.save(p("scanned-notice.png"))


def make_encrypted_pdf():
    r = PdfReader(p("kyc-guide.pdf"))
    wtr = PdfWriter()
    for page in r.pages:
        wtr.add_page(page)
    wtr.encrypt(user_password="secret", owner_password="owner")
    with open(p("protected.pdf"), "wb") as f:
        wtr.write(f)


def make_epub():
    b = epub.EpubBook()
    b.set_identifier("readit-sample-1")
    b.set_title(TITLE)
    b.set_language("en")
    ch1 = epub.EpubHtml(title="Introduction", file_name="ch1.xhtml", lang="en")
    ch1.content = f"<h1>Introduction</h1><p>{INTRO}</p>"
    ch2 = epub.EpubHtml(title="Risk", file_name="ch2.xhtml", lang="en")
    ch2.content = f"<h1>Risk rating</h1><p>{RISK}</p><ol>" + "".join(f"<li>{s}</li>" for s in STEPS) + "</ol>"
    b.add_item(ch1)
    b.add_item(ch2)
    b.toc = (ch1, ch2)
    b.add_item(epub.EpubNcx())
    b.add_item(epub.EpubNav())
    b.spine = ["nav", ch1, ch2]
    epub.write_epub(p("kyc-book.epub"), b)


def make_text_formats():
    with open(p("kyc-notes.md"), "w", encoding="utf-8") as f:
        f.write(f"# {TITLE}\n\n{INTRO}\n\n## Steps\n\n" + "".join(f"- {s}\n" for s in STEPS))
    with open(p("kyc-notes.txt"), "w", encoding="utf-8") as f:
        f.write(f"{TITLE.upper()}\n\n{INTRO}\n\n{RISK}\n")
    with open(p("kyc-page.html"), "w", encoding="utf-8") as f:
        f.write(
            f"<html><head><title>{TITLE}</title><script>alert('never run')</script></head><body>"
            f"<nav>Home | Contact</nav><h1>{TITLE}</h1><p>{INTRO}</p><footer>Footer</footer></body></html>"
        )
    bs = chr(92)
    rtf = "{" + bs + "rtf1" + bs + "ansi{" + bs + "fonttbl{" + bs + "f0 Times;}}" + bs + "pard " + TITLE + bs + "par " + RISK + bs + "par}"
    with open(p("kyc-notes.rtf"), "w", encoding="ascii") as f:
        f.write(rtf)


def make_broken():
    with open(p("broken.docx"), "wb") as f:
        f.write(b"PK\x03\x04 this is not really a zip archive")
    with open(p("unknown.xyz"), "wb") as f:
        f.write(bytes(range(256)) * 4)
    # A real zip that is not a document
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("readme.txt", "hello")
    with open(p("archive.zip"), "wb") as f:
        f.write(buf.getvalue())


if __name__ == "__main__":
    make_docx()
    make_pptx()
    make_pdf()
    make_scanned_pdf()
    make_encrypted_pdf()
    make_epub()
    make_text_formats()
    make_broken()
    print("\n".join(sorted(os.listdir(OUT))))
